from __future__ import annotations

import asyncio
import io

from docx import Document
from openpyxl import Workbook
from sqlalchemy import select

from app.connectors.llm import EmbeddingProviderError
from app.knowledge.ingest import (
    IngestionService,
    ParsedSection,
    chunk_sections,
    parse_docx,
    parse_pdf,
    parse_xlsx,
    sanitize_filename,
)
from app.knowledge.models import KnowledgeChunk, KnowledgeDocumentStatus
from app.knowledge.vector_store import EMBEDDING_DIM, InMemoryVectorStore


class FakeEmbeddingConnector:
    connector_id = "fake-embedding"
    version = "v1"

    def __init__(self, fail_with=None, fail_after_calls=None):
        self.fail_with = fail_with
        self.fail_after_calls = fail_after_calls
        self.calls: list[list[str]] = []

    async def embed(self, texts):
        self.calls.append(list(texts))
        if self.fail_with is not None and (
            self.fail_after_calls is None or len(self.calls) > self.fail_after_calls
        ):
            raise self.fail_with
        return [[1.0] * EMBEDDING_DIM for _ in texts]


class FakeStorageConnector:
    connector_id = "fake-storage"
    version = "v1"

    def __init__(self):
        self.objects: dict[str, bytes] = {}

    async def put(self, key, content):
        self.objects[key] = content

    async def get(self, key):
        return self.objects[key]

    async def delete(self, key):
        self.objects.pop(key, None)


def docx_bytes(text: str) -> bytes:
    document = Document()
    document.add_paragraph(text)
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def xlsx_bytes(rows: list[list[str]], sheet_title: str = "Products") -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = sheet_title
    for row in rows:
        sheet.append(row)
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def minimal_pdf(text: str) -> bytes:
    escaped = text.replace("\\", r"\\").replace("(", r"\(").replace(")", r"\)")
    stream = f"BT /F1 24 Tf 72 720 Td ({escaped}) Tj ET".encode("latin-1")
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        (
            b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
            b"/Contents 4 0 R /Resources << /Font << /F1 5 0 R >> >> >>"
        ),
        b"<< /Length " + str(len(stream)).encode() + b" >>\nstream\n" + stream + b"\nendstream",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    pdf = bytearray(b"%PDF-1.4\n")
    offsets: list[int] = []
    for index, obj in enumerate(objects, start=1):
        offsets.append(len(pdf))
        pdf += f"{index} 0 obj\n".encode() + obj + b"\nendobj\n"
    xref_offset = len(pdf)
    pdf += f"xref\n0 {len(objects) + 1}\n".encode()
    pdf += b"0000000000 65535 f \n"
    for offset in offsets:
        pdf += f"{offset:010d} 00000 n \n".encode()
    trailer = f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{xref_offset}\n%%EOF\n"
    pdf += trailer.encode()
    return bytes(pdf)


def make_service(session, *, embedding=None, storage=None, vector_store=None):
    return IngestionService(
        session,
        storage=storage if storage is not None else FakeStorageConnector(),
        embedding=embedding,
        vector_store=vector_store if vector_store is not None else InMemoryVectorStore(),
    )


def create_document(service, organization_id, *, content_type="application/octet-stream", size=1024):
    return service.create_document(
        organization_id=organization_id,
        filename="spec.docx",
        content_type=content_type,
        size=size,
    )


def test_create_document_starts_uploaded(session, organizations):
    service = make_service(session)

    document = create_document(service, organizations["acme"].id)

    assert document.status is KnowledgeDocumentStatus.UPLOADED


def test_process_marks_document_ready_and_persists_chunks(session, organizations):
    embedding = FakeEmbeddingConnector()
    service = make_service(session, embedding=embedding)
    document = create_document(
        service,
        organizations["acme"].id,
        content_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    )

    result = asyncio.run(
        service.process(document.id, organizations["acme"].id, docx_bytes("alpha beta gamma"))
    )

    assert result.status is KnowledgeDocumentStatus.READY
    chunks = list(session.scalars(select(KnowledgeChunk).where(KnowledgeChunk.document_id == document.id)))
    assert chunks
    assert all(chunk.organization_id == organizations["acme"].id for chunk in chunks)
    assert embedding.calls


def test_process_fails_when_embedding_is_missing(session, organizations):
    service = make_service(session, embedding=None)
    document = create_document(
        service,
        organizations["acme"].id,
        content_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    )

    result = asyncio.run(service.process(document.id, organizations["acme"].id, docx_bytes("alpha")))

    assert result.status is KnowledgeDocumentStatus.FAILED
    assert "embedding" in result.failure_message.lower()


def test_process_marks_failed_when_embedding_raises(session, organizations):
    service = make_service(
        session,
        embedding=FakeEmbeddingConnector(fail_with=EmbeddingProviderError("provider down")),
    )
    document = create_document(
        service,
        organizations["acme"].id,
        content_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    )

    result = asyncio.run(service.process(document.id, organizations["acme"].id, docx_bytes("alpha")))

    assert result.status is KnowledgeDocumentStatus.FAILED
    assert result.failure_message == "embedding provider request failed"


def test_process_rolls_back_chunks_and_vectors_when_embedding_raises_mid_way(session, organizations):
    vector_store = InMemoryVectorStore()
    embedding = FakeEmbeddingConnector(
        fail_with=EmbeddingProviderError("provider down"),
        fail_after_calls=1,
    )
    service = make_service(session, embedding=embedding, vector_store=vector_store)
    document = create_document(
        service,
        organizations["acme"].id,
        content_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    )

    result = asyncio.run(
        service.process(
            document.id,
            organizations["acme"].id,
            docx_bytes("alpha beta " * 10_000),
        )
    )

    assert result.status is KnowledgeDocumentStatus.FAILED
    assert result.failure_message == "embedding provider request failed"
    assert len(embedding.calls) > 1
    chunks = list(
        session.scalars(select(KnowledgeChunk).where(KnowledgeChunk.document_id == document.id))
    )
    assert chunks == []
    assert vector_store._records == {}


def test_process_fails_when_parsed_text_is_empty(session, organizations):
    service = make_service(session, embedding=FakeEmbeddingConnector())
    document = create_document(
        service,
        organizations["acme"].id,
        content_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    )

    result = asyncio.run(service.process(document.id, organizations["acme"].id, docx_bytes("")))

    assert result.status is KnowledgeDocumentStatus.FAILED
    assert result.failure_message == "no extractable text found in document"
    chunks = list(
        session.scalars(select(KnowledgeChunk).where(KnowledgeChunk.document_id == document.id))
    )
    assert chunks == []


def test_process_fails_on_unsupported_content_type(session, organizations):
    service = make_service(session, embedding=FakeEmbeddingConnector())
    document = create_document(service, organizations["acme"].id, content_type="application/octet-stream")

    result = asyncio.run(service.process(document.id, organizations["acme"].id, b"raw bytes"))

    assert result.status is KnowledgeDocumentStatus.FAILED
    assert "unsupported" in result.failure_message.lower()


def test_parse_docx_extracts_paragraph_text():
    sections = parse_docx(docx_bytes("Hello world"))

    assert len(sections) == 1
    assert "Hello world" in sections[0].text
    assert sections[0].page_or_sheet == ""


def test_parse_xlsx_extracts_sheet_rows():
    sections = parse_xlsx(xlsx_bytes([["SKU", "Name"], ["FL-200W", "LED Floodlight"]]))

    assert any(section.page_or_sheet == "Products" for section in sections)
    assert any("LED Floodlight" in section.text for section in sections)


def test_parse_pdf_extracts_page_text():
    sections = parse_pdf(minimal_pdf("Hello PDF"))

    assert len(sections) == 1
    assert sections[0].page_or_sheet == "page 1"
    assert "Hello PDF" in sections[0].text


def test_chunk_sections_splits_with_overlap():
    text = "word " * 100
    sections = [ParsedSection(text=text, page_or_sheet="page 1")]

    chunks = chunk_sections(sections, chunk_size=100, overlap=20)

    assert len(chunks) > 1
    assert all(chunk.token_estimate > 0 for chunk in chunks)
    assert chunks[0].page_or_sheet == "page 1"


def test_sanitize_filename_strips_paths_and_collapses_whitespace():
    assert sanitize_filename(r"C:\reports\Q3  report.docx") == "Q3 report.docx"
    assert sanitize_filename("inbox/spec.pdf") == "spec.pdf"
    assert sanitize_filename("   ") == "document"
    assert sanitize_filename("x" * 300) == "x" * 120
